import requests
from datetime import datetime, timezone, timedelta
from openpyxl import Workbook, load_workbook
import os

EXCEL_FILE = "oil_data.xlsx"
VN_TZ = timezone(timedelta(hours=7))

def get_hyperliquid():
    try:
        r = requests.post("https://api.hyperliquid.xyz/info",
                          json={"type": "metaAndAssetCtxs", "dex": "xyz"}, timeout=12)
        data = r.json()
        universe, ctxs = data[0]["universe"], data[1]
        for i, coin in enumerate(universe):
            if coin["name"] == "xyz:CL":
                ctx = ctxs[i]
                return {
                    "exchange": "Hyperliquid",
                    "pair": "xyz:CL",
                    "mark": float(ctx["markPx"]),
                    "index": float(ctx.get("oraclePx", 0)),
                    "funding": float(ctx["funding"]),
                    "oi": float(ctx.get("openInterest", 0))
                }
    except Exception as e:
        return {"exchange": "Hyperliquid", "error": str(e)}
    return {"exchange": "Hyperliquid", "error": "Not found"}

def get_vest():
    try:
        r = requests.get("https://server-prod.hz.vestmarkets.com/v2/ticker/latest?symbols=CL-PERP", timeout=12)
        t = r.json()["tickers"][0]
        return {
            "exchange": "Vest",
            "pair": "CL-PERP",
            "mark": float(t["markPrice"]),
            "index": float(t["indexPrice"]),
            "funding": float(t["oneHrFundingRate"]),
            "oi": None
        }
    except Exception as e:
        return {"exchange": "Vest", "error": str(e)}

def get_binance():
    try:
        d = requests.get("https://fapi.binance.com/fapi/v1/premiumIndex?symbol=CLUSDT", timeout=12).json()
        oi = requests.get("https://fapi.binance.com/fapi/v1/openInterest?symbol=CLUSDT", timeout=12).json()
        return {
            "exchange": "Binance",
            "pair": "CLUSDT",
            "mark": float(d["markPrice"]),
            "index": float(d["indexPrice"]),
            "funding": float(d["lastFundingRate"]),
            "oi": float(oi["openInterest"])
        }
    except Exception as e:
        return {"exchange": "Binance", "error": str(e)}

def get_ostium():
    try:
        r = requests.get("https://metadata-backend.ostium.io/PricePublish/latest-prices", timeout=12)
        data = r.json()
        items = data if isinstance(data, list) else data.get("data", [])
        for item in items:
            if item.get("from") == "CL" and item.get("to") == "USD":
                return {
                    "exchange": "Ostium",
                    "pair": "CL-USD",
                    "mark": float(item["mid"]),
                    "index": float(item["mid"]),
                    "funding": None,
                    "oi": None
                }
    except Exception as e:
        return {"exchange": "Ostium", "error": str(e)}
    return {"exchange": "Ostium", "error": "CL not found"}

def get_qfex():
    try:
        # Thử lấy refdata (có thể không có giá realtime)
        r = requests.get("https://api.qfex.com/refdata", timeout=12)
        data = r.json()
        items = data.get("data", data) if isinstance(data, dict) else data
        for item in items:
            if isinstance(item, dict) and item.get("symbol") == "CL-USD":
                price = item.get("underlier_price") or item.get("mark_price") or item.get("price")
                if price and str(price).strip():
                    return {
                        "exchange": "QFEX",
                        "pair": "CL-USD",
                        "mark": float(price),
                        "index": float(price),
                        "funding": None,
                        "oi": None
                    }
        return {"exchange": "QFEX", "error": "No live price in public API"}
    except Exception as e:
        return {"exchange": "QFEX", "error": str(e)}

def save_to_excel(rows):
    now = datetime.now(VN_TZ).strftime("%Y-%m-%d %H:%M:%S")
    if os.path.exists(EXCEL_FILE):
        wb = load_workbook(EXCEL_FILE)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.append(["Thời gian (VN)", "Sàn", "Cặp", "Mark Price", "Index Price", "Funding", "Open Interest", "Ghi chú"])

    for r in rows:
        if "error" in r:
            ws.append([now, r["exchange"], "", "", "", "", "", r["error"]])
        else:
            funding = f"{r['funding']*100:.6f}%" if r.get("funding") is not None else ""
            oi = f"{r['oi']:,.0f}" if r.get("oi") is not None else ""
            ws.append([now, r["exchange"], r["pair"], r.get("mark"), r.get("index"), funding, oi, ""])
    wb.save(EXCEL_FILE)
    print(f"✅ Đã lưu {EXCEL_FILE}")

def main():
    print("Đang lấy dữ liệu từ 5 sàn...")
    results = [
        get_hyperliquid(),
        get_vest(),
        get_binance(),
        get_ostium(),
        get_qfex()
    ]
    for r in results:
        print(r)
    save_to_excel(results)
    print("Xong!")

if __name__ == "__main__":
    main()
